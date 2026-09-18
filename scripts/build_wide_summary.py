"""
Step 4 (optional): Pivot final_output.xlsx into one-row-per-product
=======================================================================
generate_missing_images_gcp.py writes one row per product PER SLOT (long
format) — useful for auditing status per image, but not what you want to
hand to someone who just wants "the 6 images for this product" at a
glance. This pivots that into one row per product with 6 image-link
columns, each a real clickable hyperlink (not just URL-shaped text).

Usage:
    python3 build_wide_summary.py --input final_output.xlsx --out products_6_images.xlsx
"""
import argparse

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pipeline_lib import write_link_cell

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_long_format(path: str) -> tuple:
    """Returns (products dict keyed by Product_ID, ordered list of Product_IDs)."""
    src = openpyxl.load_workbook(path).active
    rows = list(src.iter_rows(min_row=2, values_only=True))

    products = {}
    order = []
    for (pid, sku, name, category, slot, image_type, status, image_source, gcp_link) in rows:
        if pid not in products:
            products[pid] = {"sku": sku, "name": name, "category": category,
                             "slots": {}, "secondary": None, "note": ""}
            order.append(pid)

        if status == "no_rule_for_category":
            products[pid]["note"] = (f"No rule for category {category!r} — cannot "
                                     f"determine required images")
            continue

        if status in ("Existing", "Existing (quality_check_failed)"):
            # image_source is still the original admin-panel URL in both
            # cases — quality_check_failed just means we couldn't verify
            # its resolution, not that it was replaced with anything.
            link, label = image_source, f"[ADMIN PHOTO] {image_source}"
        elif status == "Existing (enhanced)":
            # A real admin photo that was below the 2048px floor and got
            # upscaled — still the admin's real product photo, not an AI
            # image, so it gets its own tag distinct from "Generated"
            # rather than looking identical to one (see feedback that
            # admin-vs-AI-generated needs to be visually obvious in this
            # sheet, not just inferable from the Notes column).
            if gcp_link:
                link, label = gcp_link, f"[ADMIN PHOTO — ENHANCED TO 2K] {gcp_link}"
            else:
                link, label = None, f"[ADMIN PHOTO — ENHANCED TO 2K] {image_source} (LOCAL FILE — not uploaded to GCS)"
        elif status == "Generated":
            if gcp_link:
                link, label = gcp_link, f"[AI GENERATED] {gcp_link}"
            else:
                link, label = None, f"[AI GENERATED] {image_source} (LOCAL FILE — not uploaded to GCS)"
        elif status == "Generated (needs review)":
            # This IS a generated (uploaded) image that exhausted every
            # verify-and-retry attempt without a clean pass — a reviewer
            # needs to actually look at it to judge whether it's usable or
            # needs a manual redo, so a bare status word with nothing to
            # click isn't reviewable. Show the link, but labeled unmistakably
            # as unverified — never bare/plain like a normal passing link —
            # so nobody skims the sheet and treats it as a clean pass.
            if gcp_link:
                link, label = gcp_link, f"[NEEDS REVIEW — unverified, exhausted retries] {gcp_link}"
            else:
                link, label = None, f"{image_source} (LOCAL FILE, unverified — not uploaded to GCS)"
            note_flag = f"Slot {slot} ({image_type}) needs manual review"
            products[pid]["note"] = (products[pid]["note"] + "; " + note_flag
                                     if products[pid]["note"] else note_flag)
        elif status == "MANUAL_REVIEW_REQUIRED":
            # Not a generated deliverable at all — generation was never
            # attempted (bad/mismatched reference photo, or admin dimension
            # data that can't be trusted). image_source here is the
            # product's own REFERENCE photo (see process_slot_task), so a
            # reviewer has something real to look at — unlike the generated-
            # but-failed case above, there's no risk of mistaking a plain
            # reference photo for a finished catalog image, so the link is
            # kept, just clearly labeled as a reference, not a final image.
            #
            # ALL 6 slots share the exact same single reference photo when
            # generation never even started (nothing per-slot has diverged
            # yet), so the same URL would otherwise repeat identically 6
            # times — that reads as a bug ("why is it the same image every
            # time?") rather than the real reason (one bad photo blocks the
            # whole product). Show the full link only the first time per
            # product; later slots stay clickable (same link) but with a
            # short label pointing back to it instead of repeating it.
            seen_refs = products[pid].setdefault("_seen_ref_photos", set())
            if image_source:
                if image_source in seen_refs:
                    link, label = image_source, "[NEEDS REVIEW — same flagged reference photo, see Image 1]"
                else:
                    seen_refs.add(image_source)
                    link, label = image_source, f"[NEEDS REVIEW — reference photo, not final] {image_source}"
            else:
                link, label = None, f"[{status}]"
            note_flag = f"Slot {slot} ({image_type}) needs manual review"
            products[pid]["note"] = (products[pid]["note"] + "; " + note_flag
                                     if products[pid]["note"] else note_flag)
        elif status.startswith("Skipped"):
            # By-design, not a defect: no admin Size/Dimensions photo was
            # available and generation for this slot is disabled by policy
            # (see process_slot_task) — this product simply ships 5 images
            # instead of 6. Leave the cell blank rather than a bracketed
            # status word or a "needs review" note, since there is nothing
            # for a reviewer to act on here.
            link, label = None, ""
        else:
            link, label = None, f"[{status}]"

        # A bonus/secondary row (Slot "6b", "3b", ...) — the best-effort
        # Size/Dimensions chart generated ALONGSIDE the required alt-angle
        # primary slot, never one of the numbered 1-6 slots itself (see
        # process_slot_task's secondary_row). Goes in its own column so it
        # isn't silently dropped (a plain `slots[slot]` write would never
        # be read back — write_wide_excel only iterates integer slots 1-6).
        if isinstance(slot, str) and slot.endswith("b"):
            products[pid]["secondary"] = (label, link)
        else:
            products[pid]["slots"][slot] = (label, link)

    return products, order


def load_categories(products_csv: str) -> dict:
    """Product_ID -> (Category_L1, Category_L2, Category_L3) from
    products_detail.csv (step 1's output) — the real LO/L1/L2 category
    breakdown, not the single merged rule-category final_output.xlsx
    carries. Returns {} if the file can't be read (caller then just omits
    these columns rather than failing the whole sheet over it)."""
    try:
        df = pd.read_csv(products_csv, dtype={"Product_ID": str})
    except (OSError, pd.errors.ParserError):
        return {}
    out = {}
    for _, row in df.iterrows():
        out[row["Product_ID"]] = (
            row.get("Category_L1", "") or "",
            row.get("Category_L2", "") or "",
            row.get("Category_L3", "") or "",
        )
    return out


def write_wide_excel(products: dict, order: list, out_path: str, categories: dict = None):
    categories = categories or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.sheet_view.showGridLines = False

    headers = ["Product_ID", "SKU", "Name", "LO CATEGORY", "L1 CATEGORY", "L2 CATEGORY",
               "Image 1", "Image 2", "Image 3", "Image 4", "Image 5", "Image 6",
               "Secondary (Size/Dimensions — best effort, not guaranteed)", "Notes"]
    widths = [11, 16, 40, 14, 22, 22, 32, 32, 32, 32, 32, 32, 32, 40]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = WRAP
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"

    for r_idx, pid in enumerate(order, start=2):
        p = products[pid]
        lo, l1, l2 = categories.get(str(pid), ("", "", ""))
        plain_values = [pid, p["sku"], p["name"], lo, l1, l2]
        for col, val in enumerate(plain_values, start=1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font = BODY_FONT
            cell.alignment = WRAP
            cell.border = BORDER

        for slot_num in range(1, 7):
            label, link = p["slots"].get(slot_num, ("", None))
            write_link_cell(ws, r_idx, 6 + slot_num, label, url=link,
                            font=BODY_FONT, border=BORDER)

        sec_label, sec_link = p.get("secondary") or ("", None)
        write_link_cell(ws, r_idx, 13, sec_label, url=sec_link,
                        font=BODY_FONT, border=BORDER)

        note_cell = ws.cell(row=r_idx, column=14, value=p["note"])
        note_cell.font = BODY_FONT
        note_cell.alignment = WRAP
        note_cell.border = BORDER
        ws.row_dimensions[r_idx].height = 40

    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="final_output.xlsx from step 3")
    ap.add_argument("--products",
                    help="products_detail.csv from step 1 — if given, adds real "
                         "LO/L1/L2 category columns (omitted otherwise)")
    ap.add_argument("--out", default="products_6_images.xlsx")
    args = ap.parse_args()

    products, order = load_long_format(args.input)
    categories = load_categories(args.products) if args.products else {}
    write_wide_excel(products, order, args.out, categories)
    print(f"Done -> {args.out} ({len(order)} products)")


if __name__ == "__main__":
    main()
